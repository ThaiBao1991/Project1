# -*- coding: utf-8 -*-
import os
import openpyxl

file_path = os.path.join(os.path.dirname(__file__), 'DanhSachAnhBHCL.xlsm')

actions_map = {
    '00043': 'Là chủ bữa tiệc, tay cầm ly bia vàng bọt tuyết ngắm nghía và nghiêng đầu cười đĩnh đạc làm chủ bữa tiệc.\nSuy nghĩ: "Vui là chính, mình anh lo hết 👑"',
    '11533': 'Cười nắc nẻ ôm bụng, tay cầm ly bia giơ sang bàn 2 trêu cặp đôi bói toán Lê Cao Thanh & Mai Thảo.\nSuy nghĩ: "Xem bói chỉ tay giữa bàn nhậu, kiểu này lát nữa thầy bói Thanh bị bắt bao nguyên chầu là cái chắc!"',
    '12498': 'Ga-lăng bóc đĩa tôm nướng đầy ắp chia đều cho các chị em phụ nữ trên bàn, mồ hôi nhễ nhại nhưng cười rất tươi.\nSuy nghĩ: "Phận làm nam nhi duy nhất bàn chị em, bóc tôm mỏi tay nhưng đổi lại được chị em rót bia đầy tràn!"',
    '20023': 'Đứng dậy giơ cao ly bia bọt tuyết cười rạng rỡ như đóng quảng cáo bia Heineken.\nSuy nghĩ: "Đã lên đồ đi tiệc là phải rực rỡ nhất quán, 1-2-3 Zô hết mình rồi mai tính tiếp!"',
    '20433': 'Hai tay chống cằm tạo dáng bông hoa nở bên cạnh đĩa ốc hương, cười duyên chụp ảnh.\nSuy nghĩ: "Mặt mình góc này lên hình thon gọn quá nè, lát nhớ nhắc con Hương gửi ảnh gốc liền nha!"',
    '20474': 'Cầm đũa chỉ trỏ đĩa mực trứng chiên mắm, mắt lia khắp bàn tiệc tìm món độc lạ tiếp theo.\nSuy nghĩ: "Bia thì nhấp môi cho có không khí thôi, còn mồi ngon thì nhất định không được để sót miếng nào!"',
    '23169': 'Cầm smartphone giơ cao 45 độ, miệng liên tục hô "1, 2, 3 chu môi nào!" chụp liền 50 tấm ảnh nhóm.\nSuy nghĩ: "Chụp 50 tấm chắc chọn được 1 tấm đẹp nhất đăng Facebook câu ngàn like cho coi!"',
    '45055': 'Nghiêng đầu tạo dáng chữ V bên má, nụ cười tỏa nắng, tay cầm lon bia Saigon ướp lạnh.\nSuy nghĩ: "Uống ít mà vui nhiều, giữ gìn nhan sắc để mai còn làm hoa khôi công ty chứ lị!"',
    '49500': 'Bật camera trước quay clip TikTok bắt trọn biểu cảm ăn đùi gà của Thái Bảo và quẻ bói của thầy Thanh.\nSuy nghĩ: "Quay lại quả clip đắt giá này, mai đăng lên nhóm Zalo công ty thì bao cười cả tuần!"',
    '12690': 'Cầm micro bằng cả hai tay, mắt nhắm nghiền phiêu nốt cao bài hit bolero chào mừng đại tiệc.\nSuy nghĩ: "Đam mê ca hát cháy bỏng từ nhỏ, hôm nay có bia vào giọng mình ấm như ca sĩ Đan Trường vậy!"',
    '52270': 'Cầm hai chiếc khăn giấy trắng vẫy vẫy múa phụ họa uốn lượn theo nhịp hát của ca sĩ Hiếu.\nSuy nghĩ: "Làm fan cuồng cổ vũ nhiệt tình thế này, lát nữa anh Hiếu phải mời mình một ly bia cảm ơn nha!"',
    '52401': 'Cầm hai chiếc đũa gõ nhịp chan chát vào miệng bát làm trống phụ họa cho ban nhạc bàn 5.\nSuy nghĩ: "Tùng cắc tùng cắc, beat này hơi bị cháy, bàn mình quẩy xôm nhất cái quán nhậu này rồi!"',
    '10863': 'Đứng tựa mép bàn VIP, ngoái cổ sang bàn số 4 bên cạnh, giơ ly bia vẫy vẫy mời chào rôm rả.\nSuy nghĩ: "Phải kéo nhóm 4 sang cụng ly với Giám đốc Bình cho không khí nóng lên mới được!"',
    '10883': 'Bưng đĩa tôm mực nướng bốc khói nghi ngút tiếp mồi cho các sếp bàn VIP, tay kia làm động tác bắn tim.\nSuy nghĩ: "Mồi bàn VIP vơi nhanh quá, mình đi tiếp tế thế này các sếp khen chu đáo là cái chắc!"',
    '11466': 'Ngồi chống cằm một tay tạo dáng Người suy tư, mắt nhìn xa xăm giữa tiếng hò reo zô ầm ĩ.\nSuy nghĩ: "Uống bia thế này tối về ngứa kiếm ai gãi đây...? 🤔"',
    '20301': 'Tay cầm đũa gắp miếng mực nướng đưa lên miệng nhưng dừng lại giữa chừng, mắt chớp chớp suy tư.\nSuy nghĩ: "Không biết chi nhánh sữa chua thứ 2 của mình hôm nay bán được mấy trăm hộp nhỉ, tí phải mở app kiểm tra doanh thu mới yên tâm nhậu tiếp!"',
    '24040': 'Lắc lư theo điệu nhạc, một tay cầm ly bia sóng sánh, một tay giơ ngón tay cái cổ vũ bàn 5.\nSuy nghĩ: "Nhạc hay bia ngon bạn hiền, tinh thần công ty hôm nay đúng là đỉnh nóc kịch trần bay phấp phới!"',
    '20122': 'Vỗ tay đều đặn theo nhịp zô, nụ cười rạng rỡ khuấy động phong trào bàn chị em.\nSuy nghĩ: "Cứ vỗ tay cổ vũ thật to, lát nữa mấy anh bàn VIP qua mời là chị em mình không ngán ly nào!"',
    '20709': 'Cầm vại bia cụng thật kêu với người bên cạnh, bọt bia bắn tung tóe đầy phấn khích.\nSuy nghĩ: "Đã cầm ly lên là phải uống cạn đáy cốc, phong cách của Sau là không bao giờ để bia long đền!"',
    '23866': 'Nụ cười má lúm đồng tiền duyên dáng, tay bưng đĩa dưa hấu tráng miệng mời cả bàn giải nhiệt.\nSuy nghĩ: "Uống bia xong ăn miếng dưa hấu mát lịm này thì tỉnh táo ngay, lát còn quẩy hiệp 2 tăng 3!"',
    '34873': 'Đang nhanh nhẹn phân phát khăn ướt và gắp đá viên bỏ vào ly bia cho từng người trong bàn.\nSuy nghĩ: "Hậu cần chu đáo số 1 là đây, bia phải mát lạnh uống mới đã cái nư!"',
    '10630': 'Đứng bật dậy một chân gác nhẹ lên ghế, một tay giơ cao chai bia hô vang: "1 - 2 - 3... ZÔ 100%!"\nSuy nghĩ: "Khí thế hôm nay phải ngút trời, anh em phòng mình nhậu là phải nhất nhì công ty!"',
    '10091': 'Hai tay nâng vại bia to bự hướng về Giám đốc Bình, người hơi cúi cung kính, nụ cười hết nấc.\nSuy nghĩ: "Giám đốc Bình nói vậy mời Giám đốc 1 ly! 🍻"',
    '51727': 'Mắt mở to tròn xoe, tay che miệng há hốc kinh ngạc khi đưa tay cho thầy bói Lê Cao Thanh xem.\nSuy nghĩ: "Trời ơi thầy Thanh phán chuẩn dữ dội, hồi chiều vừa lén mua sắm online hết nửa tháng lương!"',
    '52795': 'Cầm ly bia quay sang cụng giao lưu với bàn VIP, nụ cười tươi tắn tự tin.\nSuy nghĩ: "Giao lưu với sếp một ly chúc công ty năm nay thưởng Tết gấp ba cho anh em phấn khởi!"',
    '12953': 'Hai mắt hình trái tim lấp lánh, miệng chảy nước miếng, tay ôm chặt chiếc đùi gà nướng bự chảng.\nSuy nghĩ: "Mồi ngon hè... Mấy ông cứ tranh nhau uống đi, một mình tôi ôm trọn cái đùi gà này là ấm bụng nhất trần đời!"',
    '11481': 'Đứng nghiêng người rót bia điêu luyện từ trên cao xuống ly, canh đúng chuẩn tỉ lệ 7 phần bia 3 phần bọt.\nSuy nghĩ: "Nghệ thuật rót bia đỉnh cao này phải luyện chục năm mới được, bọt mịn như kem tuyết!"',
    '12333': 'Cầm ly bia hiên ngang bước sang bàn VIP, hai tay nâng ly kính cẩn chúc mừng Giám đốc Bình và hai sếp Vũ, Tân.\nSuy nghĩ: "Cơ hội vàng ngàn năm có một để thể hiện bản lĩnh tửu lượng trước mặt Giám đốc Bình và các Sếp!"',
    '13011': 'Đeo kính râm đen, mắt lim dim làm bộ cao nhân đắc đạo, nắm chặt tay Mai Thảo phán chỉ tay thần sầu.\nSuy nghĩ: "Nhìn bàn tay này là biết số hưởng rồi, nhưng lát nữa tiền tính bill bàn này thì chắc chắn là em chịu nha Thảo!"',
    '20434': 'Nữ tướng oai phong, một tay chống hông, một tay nâng vại bia không hề nao núng trước phái nam.\nSuy nghĩ: "Đừng nghĩ chị em phụ nữ không biết uống, bàn 6 hôm nay tôi chấp hết các đấng mày râu!"',
    '13038': 'Tay khoác vai đồng đội, tay kia giơ nắm đấm hô hào quyết tâm chiến đấu tới giọt bia cuối cùng.\nSuy nghĩ: "Đã ngồi bàn 6 là tinh thần thép, không say không về, ai trốn về sớm phạt 1 két!"',
    '13044': 'Ngồi đếm hàng vỏ chai bia rỗng xếp dài dưới chân bàn với nụ cười cực kỳ đắc chí và tự hào.\nSuy nghĩ: "Mới xếp được có 3 hàng vỏ chai thôi mà, tối nay phải lập kỷ lục Guinness quán nhậu này mới chịu!"',
    '13075': 'Lấy vạt áo lau mồ hôi trên trán cười toe toét, mắt vẫn sáng rực đầy năng lượng chiến đấu.\nSuy nghĩ: "Nóng quá nhưng mà vui quá trời! Em vẫn còn tỉnh queo nha anh Long ơi, rót tiếp đi!"'
}

print(f"Loading {file_path}...")
wb = openpyxl.load_workbook(file_path, keep_vba=True)
ws = wb.active

updated_count = 0
for r in range(2, ws.max_row + 1):
    mnv = str(ws.cell(r, 2).value or '').strip()
    if mnv in actions_map:
        ws.cell(r, 5).value = actions_map[mnv]
        updated_count += 1
        print(f"Row {r} (MNV {mnv}): Updated")

wb.save(file_path)
print(f"SUCCESS: Updated all {updated_count} rows in {file_path}!")
